"""
parser.py - Text & Feature Extraction Engine for Resume Analyzer.
Extracts text from PDF/DOCX files, parses key candidate details, contact info,
identifies resume structural sections, and extracts skills using spaCy and an extensive skill taxonomy.
"""

import io
import re
import fitz  # PyMuPDF
import docx
import spacy
from typing import Dict, List, Tuple, Any, Set

# Extensive Predefined Skill Taxonomy categorized for deep analysis
SKILL_TAXONOMY: Dict[str, List[str]] = {
    "Programming Languages": [
        "Python", "Java", "C++", "C#", "C", "JavaScript", "TypeScript", "Go", "Golang",
        "Rust", "Ruby", "PHP", "Swift", "Kotlin", "R", "SQL", "PL/SQL", "T-SQL",
        "HTML", "HTML5", "CSS", "CSS3", "Sass", "Bash", "Shell", "Powershell", "Scala", "MATLAB", "Perl", "Dart"
    ],
    "Frameworks & Libraries": [
        "React", "React.js", "Angular", "Vue", "Vue.js", "Node.js", "Express.js", "Django", "Flask",
        "FastAPI", "Spring Boot", "Spring", "Next.js", "Nuxt.js", "NestJS", "Tailwind CSS",
        "Bootstrap", "jQuery", "Redux", "GraphQL", "PyTorch", "TensorFlow", "Keras",
        "Scikit-learn", "Pandas", "NumPy", "SciPy", "OpenCV", "NLTK", "spaCy", "Hugging Face",
        "Transformers", "LangChain", "LlamaIndex", "Streamlit", "Gradio", "Spark", "PySpark",
        "Hadoop", "RxJS", "ASP.NET", ".NET Core", "Laravel", "Symfony"
    ],
    "Databases & Data Stores": [
        "PostgreSQL", "MySQL", "MongoDB", "Redis", "SQLite", "Oracle", "Microsoft SQL Server",
        "Cassandra", "DynamoDB", "Neo4j", "Elasticsearch", "Snowflake", "BigQuery",
        "Firebase", "Supabase", "CockroachDB", "MariaDB", "CouchDB", "Vector Database",
        "Pinecone", "Weaviate", "Qdrant", "ChromaDB", "Milvus"
    ],
    "Cloud & DevOps": [
        "AWS", "Amazon Web Services", "Azure", "Microsoft Azure", "GCP", "Google Cloud Platform",
        "Docker", "Kubernetes", "K8s", "Terraform", "Ansible", "Jenkins", "GitHub Actions",
        "GitLab CI", "CI/CD", "Linux", "Unix", "Ubuntu", "Nginx", "Apache", "Serverless",
        "AWS Lambda", "Cloudflare", "Helm", "Prometheus", "Grafana", "ArgoCD", "CloudFormation"
    ],
    "AI / ML & Data Science": [
        "Machine Learning", "Deep Learning", "Natural Language Processing", "NLP",
        "Computer Vision", "Generative AI", "LLM", "LLMs", "Large Language Models",
        "Prompt Engineering", "MLOps", "Data Engineering", "Data Visualization",
        "Reinforcement Learning", "Neural Networks", "CNN", "RNN", "LSTM", "Transformers",
        "BERT", "GPT", "GANs", "Feature Engineering", "A/B Testing", "Statistical Analysis"
    ],
    "Tools & Methodologies": [
        "Git", "GitHub", "GitLab", "Bitbucket", "Jira", "Confluence", "Agile", "Scrum",
        "Kanban", "REST API", "RESTful APIs", "Microservices", "System Design",
        "Unit Testing", "Integration Testing", "TDD", "BDD", "Postman", "Swagger",
        "Webpack", "Vite", "Maven", "Gradle", "NPM", "Yarn", "VS Code"
    ],
    "Soft Skills & Leadership": [
        "Communication", "Leadership", "Problem Solving", "Critical Thinking",
        "Teamwork", "Collaboration", "Time Management", "Adaptability", "Agile Mindset",
        "Project Management", "Strategic Planning", "Negotiation", "Cross-functional Team Leadership",
        "Mentorship", "Stakeholder Management", "Decision Making", "Analytical Thinking"
    ]
}

# Pre-process a flattened list of canonical skill names and regex patterns
def _build_skill_patterns() -> List[Tuple[str, re.Pattern]]:
    patterns = []
    seen = set()
    for category, skills in SKILL_TAXONOMY.items():
        for skill in skills:
            if skill.lower() in seen:
                continue
            seen.add(skill.lower())
            
            # Escape regex special characters, handle special cases like C++, C#, .NET, Node.js
            escaped_skill = re.escape(skill)
            # Ensure word boundaries or non-alphanumeric boundaries
            pattern_str = r'(?<![A-Za-z0-9_#+.-])' + escaped_skill + r'(?![A-Za-z0-9_#+.-])'
            try:
                pattern = re.compile(pattern_str, re.IGNORECASE)
                patterns.append((skill, pattern))
            except re.error:
                continue
    return patterns

ALL_SKILL_PATTERNS = _build_skill_patterns()

# Lazy loading spaCy model to prevent unnecessary overhead
_nlp_model = None

def get_spacy_model():
    """Loads and caches spaCy NLP model safely with fallback."""
    global _nlp_model
    if _nlp_model is None:
        try:
            _nlp_model = spacy.load("en_core_web_sm")
        except Exception:
            try:
                # Attempt blank English pipeline fallback if model not pre-downloaded
                _nlp_model = spacy.blank("en")
            except Exception as e:
                raise RuntimeError(f"Failed to load spaCy model: {e}")
    return _nlp_model


def extract_text_from_pdf(file_input: Any) -> str:
    """Extracts raw text from a PDF file (stream, bytes, or file path)."""
    text = ""
    try:
        if isinstance(file_input, (bytes, bytearray)):
            doc = fitz.open(stream=file_input, filetype="pdf")
        elif hasattr(file_input, "read"):
            file_bytes = file_input.read()
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        else:
            doc = fitz.open(str(file_input))
            
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                text += page_text + "\n"
        doc.close()
    except Exception as e:
        raise ValueError(f"Error reading PDF file: {e}")
    return text.strip()


def extract_text_from_docx(file_input: Any) -> str:
    """Extracts raw text from a DOCX file (stream, bytes, or file path)."""
    text = ""
    try:
        if isinstance(file_input, (bytes, bytearray)):
            doc = docx.Document(io.BytesIO(file_input))
        elif hasattr(file_input, "read"):
            file_bytes = file_input.read()
            doc = docx.Document(io.BytesIO(file_bytes))
        else:
            doc = docx.Document(str(file_input))
            
        # Extract from paragraphs
        for p in doc.paragraphs:
            if p.text:
                text += p.text + "\n"
                
        # Extract from tables (resumes often use tables for layout)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text += row_text + "\n"
    except Exception as e:
        raise ValueError(f"Error reading DOCX file: {e}")
    return text.strip()


def parse_resume_file(uploaded_file: Any) -> Tuple[str, str]:
    """
    Reads an uploaded file object or path and returns (extracted_text, file_type).
    """
    file_name = getattr(uploaded_file, "name", str(uploaded_file)).lower()
    
    if file_name.endswith(".pdf"):
        text = extract_text_from_pdf(uploaded_file)
        file_type = "PDF"
    elif file_name.endswith(".docx") or file_name.endswith(".doc"):
        text = extract_text_from_docx(uploaded_file)
        file_type = "DOCX"
    else:
        raise ValueError("Unsupported file format. Please upload a PDF or DOCX resume.")
        
    if not text.strip():
        raise ValueError("Could not extract any readable text from the file. Ensure it is not a scanned image PDF.")
        
    return text, file_type


def extract_contact_info(text: str) -> Dict[str, str]:
    """Extracts email and phone number using regular expressions."""
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    phone_pattern = r'(\+?\d{1,3}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}'
    
    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)
    
    return {
        "email": emails[0] if emails else "Not Found",
        "phone": phones[0] if phones else "Not Found"
    }


def extract_sections(text: str) -> Dict[str, bool]:
    """
    Checks for the presence of major structural sections required in an ATS-friendly resume.
    """
    text_lower = text.lower()
    sections = {
        "Work Experience": bool(re.search(r'\b(experience|work history|employment|career|work experience)\b', text_lower)),
        "Education": bool(re.search(r'\b(education|academic|qualifications|degree|university)\b', text_lower)),
        "Projects": bool(re.search(r'\b(projects|personal projects|key projects|portfolio)\b', text_lower)),
        "Skills": bool(re.search(r'\b(skills|technical skills|core competencies|technologies|proficiencies)\b', text_lower)),
        "Certifications": bool(re.search(r'\b(certifications|licenses|courses|certificates|achievements)\b', text_lower))
    }
    return sections


def extract_skills_from_text(text: str) -> Dict[str, Any]:
    """
    Matches text against the Skill Taxonomy and returns structured skill findings:
    - flat_skills: Set of unique skill names matched
    - categorized_skills: Dict mapping category -> list of extracted skills
    """
    extracted_flat: Set[str] = set()
    categorized: Dict[str, List[str]] = {cat: [] for cat in SKILL_TAXONOMY.keys()}
    
    # 1. Regex taxonomy matching
    for skill_name, pattern in ALL_SKILL_PATTERNS:
        if pattern.search(text):
            extracted_flat.add(skill_name)
            # Find category
            for category, skill_list in SKILL_TAXONOMY.items():
                if any(s.lower() == skill_name.lower() for s in skill_list):
                    if skill_name not in categorized[category]:
                        categorized[category].append(skill_name)
                        
    # 2. spaCy Named Entity Recognition for extra entities (e.g. ORG, PRODUCT)
    nlp = get_spacy_model()
    try:
        doc = nlp(text[:100000])  # Cap long text for safety
        for ent in doc.ents:
            if ent.label_ in ["ORG", "PRODUCT"] and len(ent.text) < 30:
                ent_text = ent.text.strip()
                # If entity matches known skills
                for skill_name, _ in ALL_SKILL_PATTERNS:
                    if skill_name.lower() == ent_text.lower():
                        extracted_flat.add(skill_name)
    except Exception:
        pass
        
    # Remove empty categories
    categorized = {k: v for k, v in categorized.items() if len(v) > 0}
    
    return {
        "flat_skills": sorted(list(extracted_flat)),
        "categorized_skills": categorized,
        "total_count": len(extracted_flat)
    }


def parse_resume_content(uploaded_file: Any) -> Dict[str, Any]:
    """
    Main entry point for resume parsing.
    Returns clean parsed resume dictionary.
    """
    raw_text, file_type = parse_resume_file(uploaded_file)
    contact_info = extract_contact_info(raw_text)
    sections = extract_sections(raw_text)
    skill_data = extract_skills_from_text(raw_text)
    
    return {
        "raw_text": raw_text,
        "file_type": file_type,
        "contact_info": contact_info,
        "sections": sections,
        "skills": skill_data["flat_skills"],
        "categorized_skills": skill_data["categorized_skills"],
        "skill_count": skill_data["total_count"],
        "word_count": len(raw_text.split())
    }
